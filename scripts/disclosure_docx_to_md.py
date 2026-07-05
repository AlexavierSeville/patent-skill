#!/usr/bin/env python3
"""Convert a disclosure DOCX to Markdown while preserving highlights and comments.

Highlighted runs are wrapped as ==text==. Word comments anchored to a
paragraph are emitted immediately after that paragraph as blockquotes.
"""

from __future__ import annotations

import argparse
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def load_comments(docx_path: Path) -> dict[str, dict[str, str]]:
    try:
        with zipfile.ZipFile(str(docx_path)) as archive:
            if "word/comments.xml" not in archive.namelist():
                return {}
            xml = archive.read("word/comments.xml")
    except (KeyError, zipfile.BadZipFile):
        return {}

    comments: dict[str, dict[str, str]] = {}
    root = ET.fromstring(xml)
    for comment in root.findall(".//w:comment", NS):
        comment_id = comment.get(f"{{{W_NS}}}id")
        if comment_id is None:
            continue
        author = comment.get(f"{{{W_NS}}}author") or ""
        text = "".join(node.text or "" for node in comment.findall(".//w:t", NS)).strip()
        comments[comment_id] = {"author": author, "text": text}
    return comments


def paragraph_comment_ids(paragraph) -> list[str]:
    ids: list[str] = []
    for elem in paragraph._element.iter():
        if elem.tag in (qn("w:commentRangeStart"), qn("w:commentReference")):
            comment_id = elem.get(qn("w:id"))
            if comment_id is not None and comment_id not in ids:
                ids.append(comment_id)
    return ids


def paragraph_text_with_highlights(paragraph) -> str:
    parts: list[str] = []
    for run in paragraph.runs:
        if not run.text:
            continue
        if run.font.highlight_color is not None:
            parts.append(f"=={run.text}==")
        else:
            parts.append(run.text)
    return "".join(parts)


def markdown_line_for_paragraph(paragraph, text: str) -> str:
    style_name = paragraph.style.name if paragraph.style is not None else ""
    if style_name == "Heading 1":
        return f"# {text}" if text.strip() else ""
    if style_name.startswith("Heading"):
        digits = "".join(ch for ch in style_name if ch.isdigit())
        level = int(digits or "2")
        return f"{'#' * level} {text}" if text.strip() else ""
    return text


def convert(docx_path: Path, md_path: Path) -> dict[str, int]:
    doc = Document(str(docx_path))
    comments = load_comments(docx_path)
    seen_comments: set[str] = set()
    highlights_count = 0
    output: list[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph_text_with_highlights(paragraph)
        highlights_count += text.count("==") // 2
        output.append(markdown_line_for_paragraph(paragraph, text))

        for comment_id in paragraph_comment_ids(paragraph):
            if comment_id in seen_comments:
                continue
            seen_comments.add(comment_id)
            comment = comments.get(comment_id)
            if not comment:
                continue
            output.append(f"> 💡 [批注 #{comment_id} by {comment['author']}]: {comment['text']}")

    md_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text("\n".join(output), encoding="utf-8")
    return {
        "paragraphs": len(doc.paragraphs),
        "comments_total": len(comments),
        "comments_anchored": len(seen_comments),
        "highlights_marked": highlights_count,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Convert disclosure DOCX to annotated Markdown")
    parser.add_argument("--input", required=True, type=Path, help="source disclosure DOCX")
    parser.add_argument("--output", required=True, type=Path, help="output Markdown path")
    args = parser.parse_args(argv)

    if not args.input.is_file():
        print(f"error: input not found: {args.input}", file=sys.stderr)
        return 2

    stats = convert(args.input, args.output)
    print(
        f"wrote {args.output} "
        f"(paragraphs={stats['paragraphs']}, "
        f"comments={stats['comments_anchored']}/{stats['comments_total']}, "
        f"highlights={stats['highlights_marked']})"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
