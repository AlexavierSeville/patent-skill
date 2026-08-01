from pathlib import Path
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_COLOR_INDEX


SKILL_DIR = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = SKILL_DIR / "scripts"


def run_script(script_name, *args):
    return subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / script_name), *map(str, args)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=True,
    )


def all_suspects(data):
    """把 suspect_manifest 摊平成扁平列表 (每条附 owner).

    脚本 JSON 不再输出顶层 `suspects`(与 manifest 同源、重复占比 26.6%),
    suspect 的唯一出口是 suspect_manifest.<owner>。测试统一经本函数读取。
    """
    out = []
    for owner, items in (data.get("suspect_manifest") or {}).items():
        for it in items:
            out.append({**it, "owner": owner})
    return out


def run_script_allow_fail(script_name, *args):
    return subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / script_name), *map(str, args)],
        text=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )


GOOD_FULL_DRAFT_MD = """## 权利要求书

1.一种测试对象控制方法，其特征在于，包括：获取输入数据；根据输入数据确定中间结果；根据中间结果生成控制信号。

2.根据权利要求1所述的测试对象控制方法，其特征在于，所述获取输入数据，包括：读取传感器数据；对传感器数据滤波。

3.一种测试对象控制系统，包括存储器、处理器，其特征在于，所述处理器执行程序时实现权利要求1至2任一项所述的方法的步骤。

## 技术领域

本发明涉及测试领域。

## 背景技术

现有技术存在问题。

## 发明内容

本发明提供一种测试对象控制方法。

## 附图说明

图1为本发明实施例提供的测试对象控制方法流程示意图；

图2为本发明实施例提供的测试对象控制系统结构示意图。

## 具体实施方式

在步骤S11中，获取输入数据，包括：读取传感器数据；对传感器数据滤波。

在步骤S12中，根据输入数据确定中间结果。

在步骤S13中，根据中间结果生成控制信号。
"""


def add_minimal_comment(docx_path: Path, comment_text: str):
    tmp_path = docx_path.with_suffix(".tmp.docx")
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    ET.register_namespace("w", ns["w"])

    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        root = ET.fromstring(zin.read("word/document.xml"))
        paragraph = root.find(".//w:p", ns)
        run = paragraph.find("w:r", ns)
        paragraph.insert(0, ET.Element(f"{{{ns['w']}}}commentRangeStart", {f"{{{ns['w']}}}id": "0"}))
        paragraph.insert(list(paragraph).index(run) + 1, ET.Element(f"{{{ns['w']}}}commentRangeEnd", {f"{{{ns['w']}}}id": "0"}))
        comment_ref_run = ET.Element(f"{{{ns['w']}}}r")
        comment_ref = ET.SubElement(comment_ref_run, f"{{{ns['w']}}}commentReference")
        comment_ref.set(f"{{{ns['w']}}}id", "0")
        paragraph.append(comment_ref_run)

        comments = ET.Element(f"{{{ns['w']}}}comments")
        comment = ET.SubElement(
            comments,
            f"{{{ns['w']}}}comment",
            {f"{{{ns['w']}}}id": "0", f"{{{ns['w']}}}author": "tester"},
        )
        p = ET.SubElement(comment, f"{{{ns['w']}}}p")
        r = ET.SubElement(p, f"{{{ns['w']}}}r")
        t = ET.SubElement(r, f"{{{ns['w']}}}t")
        t.text = comment_text

        for item in zin.infolist():
            if item.filename == "word/document.xml":
                zout.writestr(item, ET.tostring(root, encoding="utf-8", xml_declaration=True))
            else:
                zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/comments.xml", ET.tostring(comments, encoding="utf-8", xml_declaration=True))

    shutil.move(tmp_path, docx_path)


def _find_pandoc():
    found = shutil.which("pandoc")
    if found:
        return found
    for cand in (
        Path.home() / ".local/bin/pandoc",
        Path.home() / "miniconda3/bin/pandoc",
        Path("/opt/homebrew/bin/pandoc"),
        Path("/usr/local/bin/pandoc"),
    ):
        if cand.exists():
            return str(cand)
    return None


PANDOC_AVAILABLE = _find_pandoc() is not None


class PatentScriptSmokeTests(unittest.TestCase):
    @unittest.skipUnless(PANDOC_AVAILABLE, "pandoc 不可用, 跳过公式 OMML 生成测试")
    def test_omml_formulas_gen_produces_block_omml(self):
        import json

        result = run_script(
            "omml_formulas.py", "gen",
            "--latex", r"Z_k=\mathrm{LN}(Z_{k-1}+\mathrm{ReLU}(x))",
            "--number", "12",
        )
        data = json.loads(result.stdout)
        self.assertEqual(len(data), 1)
        xml = data[0]["xml"]
        # 块级 oMathPara + 编号内嵌, WPS 才能渲染为可编辑二维公式 (G8-3)
        self.assertIn("m:oMathPara", xml)
        self.assertIn("(12)", xml)
        # pStyle 已剥离 (不依赖目标文档样式表); 无段落级居中 (由 oMathParaPr 承担)
        self.assertNotIn("pStyle", xml)
        self.assertNotIn('w:jc w:val="center"', xml)
        self.assertEqual(data[0]["m_d_count"], 0)

    @unittest.skipUnless(PANDOC_AVAILABLE, "pandoc 不可用, 跳过公式 OMML 生成测试")
    def test_omml_formulas_gen_normalizes_left_right(self):
        import json

        # \left( \right) 生成 <m:d> 可伸缩定界符, WPS 深嵌套渲染留白 (C-DOCX-8);
        # gen 默认归一为普通括号
        result = run_script(
            "omml_formulas.py", "gen",
            "--latex", r"y=\left(\frac{a}{b}\right)+\left|c\right|",
        )
        data = json.loads(result.stdout)
        self.assertEqual(data[0]["m_d_count"], 0)
        self.assertNotIn("<m:d>", data[0]["xml"])

    def test_omml_formulas_fix_settings_and_check(self):
        import json

        with tempfile.TemporaryDirectory() as tmp:
            docx_path = Path(tmp) / "math.docx"
            doc = Document()
            doc.add_paragraph("正文")
            doc.save(docx_path)

            # python-docx 默认模板 mathPr 带 Cambria Math + defJc/dispDef 等降级源
            run_script("omml_formulas.py", "fix-settings", str(docx_path), "--font", "STIX Two Math")
            result = run_script("omml_formulas.py", "check", str(docx_path))
            data = json.loads(result.stdout)
            self.assertTrue(data["mathpr_simplified"])
            self.assertEqual(data["math_font"], "STIX Two Math")
            self.assertEqual(data["empty_shell_count"], 0)
            self.assertEqual(data["errors"], [])

    def test_native_formula_workflow_is_wired(self):
        global_text = (SKILL_DIR / "references" / "rules" / "global.md").read_text(encoding="utf-8")
        docx_text = (SKILL_DIR / "references" / "rules" / "docx-template.md").read_text(encoding="utf-8")
        case_text = (SKILL_DIR / "references" / "cases" / "docx-execution.md").read_text(encoding="utf-8")

        # G6-1 落盘策略: md 层仍写 LaTeX 源码, DOCX 层转原生 OMML
        self.assertIn("omml_formulas.py", global_text)
        self.assertIn("原生可编辑二维公式", global_text)
        # 执行层细则 G8-3 与案例 C-DOCX-8 存在且互相指向
        self.assertIn("G8-3", docx_text)
        self.assertIn("omml_formulas.py", docx_text)
        self.assertIn("C-DOCX-8", docx_text)
        self.assertIn("C-DOCX-8", case_text)
        self.assertIn("幽灵字体", case_text)
        self.assertTrue((SCRIPTS_DIR / "omml_formulas.py").exists())
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        self.assertIn("omml_formulas.py", skill_text)

    def test_bundled_template_has_no_ghost_font(self):
        # 模板样式链引用系统不存在的字体 (如 Dutch801 Rm BT) 会让 WPS 公式
        # 正体部分空白 (C-DOCX-8 实测根因), 模板资产必须保持无幽灵字体
        template = SKILL_DIR / "assets" / "docx" / "专利撰写模板.docx"
        with zipfile.ZipFile(template) as z:
            styles = z.read("word/styles.xml").decode("utf-8")
        self.assertNotIn("Dutch801", styles)
        self.assertIn("Times New Roman", styles)

    def test_disclosure_docx_to_md_preserves_highlights_and_comments(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            docx_path = tmp_path / "disclosure.docx"
            md_path = tmp_path / "disclosure_annotated.md"

            doc = Document()
            para = doc.add_paragraph()
            para.add_run("普通内容 ")
            highlighted = para.add_run("关键创新步骤")
            highlighted.font.highlight_color = WD_COLOR_INDEX.YELLOW
            doc.save(docx_path)
            add_minimal_comment(docx_path, "该步骤必须进入权利要求1")

            result = run_script("disclosure_docx_to_md.py", "--input", docx_path, "--output", md_path)

            text = md_path.read_text(encoding="utf-8")
            self.assertIn("==关键创新步骤==", text)
            self.assertIn("批注 #0", text)
            self.assertIn("该步骤必须进入权利要求1", text)
            self.assertIn("highlights=1", result.stdout)
            self.assertIn("comments=1/1", result.stdout)

    def test_skill_frontmatter_and_bundled_template_paths_are_stable(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")

        self.assertIn("name: patent", skill_text)
        self.assertIn("description: Use when", skill_text)
        self.assertTrue((SKILL_DIR / "assets" / "docx" / "专利撰写模板.docx").exists())
        self.assertIn("assets/docx/专利撰写模板.docx", skill_text)
        self.assertNotIn("assets/专利撰写模板.docx", skill_text)

    def test_claims_format_guidance_is_bundled_not_external_path(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        claims_text = (SKILL_DIR / "references" / "rules" / "claims.md").read_text(encoding="utf-8")
        reference_path = SKILL_DIR / "references" / "cases" / "claims-format-standard.md"

        self.assertTrue(reference_path.exists())
        self.assertIn("references/cases/claims-format-standard.md", skill_text)
        self.assertIn("references/cases/claims-format-standard.md", claims_text)
        self.assertNotIn("/Users/nafsae/Desktop/Patent/第3-4课的案例资料", skill_text)
        self.assertNotIn("/Users/nafsae/Desktop/Patent/第3-4课的案例资料", claims_text)

    def test_rules_have_no_self_referential_or_docx_boundary_residue(self):
        rules_text = (SKILL_DIR / "rules.md").read_text(encoding="utf-8")
        docx_text = (SKILL_DIR / "references" / "rules" / "docx-template.md").read_text(encoding="utf-8")

        self.assertNotIn("本文件是 `rules.md` 的分级重组与扩展版", rules_text)
        self.assertIn("当前唯一规则索引", rules_text)
        self.assertNotIn("可用 zipfile + xml.etree 直接读写", docx_text)
        self.assertIn("已显式加载官方 `docx` skill", docx_text)
        self.assertIn("不得替代 `docx` skill 的执行层", docx_text)
    def test_progressive_rules_files_exist_and_are_referenced(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        rules_text = (SKILL_DIR / "rules.md").read_text(encoding="utf-8")
        rule_files = [
            "global.md",
            "claims.md",
            "full-draft.md",
            "revision.md",
            "docx-template.md",
            "figures.md",
        ]

        for filename in rule_files:
            relative_path = f"references/rules/{filename}"
            self.assertTrue((SKILL_DIR / relative_path).exists(), relative_path)
            self.assertIn(relative_path, rules_text)

        self.assertIn("按阶段读取规则", skill_text)
        self.assertIn("references/rules/global.md", skill_text)

    def test_rules_index_is_lightweight(self):
        rules_text = (SKILL_DIR / "rules.md").read_text(encoding="utf-8")

        self.assertLess(len(rules_text.splitlines()), 180)
        self.assertIn("当前唯一规则索引", rules_text)
        self.assertNotIn("# 第一部分：全局规则", rules_text)
        self.assertNotIn("# 第二部分：局部规则", rules_text)
        self.assertNotIn("# 第三部分：最终交付前总自检", rules_text)

    def test_skill_uses_progressive_rule_loading(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")

        self.assertNotIn("必须阅读本目录下的 `rules.md`。它是本工作流的强制写作规范。", skill_text)
        self.assertIn("references/rules/global.md", skill_text)
        self.assertIn("references/rules/claims.md", skill_text)
        self.assertIn("references/rules/full-draft.md", skill_text)
        self.assertIn("references/rules/revision.md", skill_text)
        self.assertIn("references/rules/docx-template.md", skill_text)
        self.assertIn("references/rules/figures.md", skill_text)

    def test_split_rules_preserve_critical_constraints(self):
        global_text = (SKILL_DIR / "references" / "rules" / "global.md").read_text(encoding="utf-8")
        claims_text = (SKILL_DIR / "references" / "rules" / "claims.md").read_text(encoding="utf-8")
        full_text = (SKILL_DIR / "references" / "rules" / "full-draft.md").read_text(encoding="utf-8")
        revision_text = (SKILL_DIR / "references" / "rules" / "revision.md").read_text(encoding="utf-8")
        docx_text = (SKILL_DIR / "references" / "rules" / "docx-template.md").read_text(encoding="utf-8")
        figures_text = (SKILL_DIR / "references" / "rules" / "figures.md").read_text(encoding="utf-8")

        self.assertIn("技术链条闭合", global_text)
        self.assertIn("禁用或风险措辞", global_text)
        self.assertIn("公式、模型、阈值", global_text)

        self.assertIn("权利要求书规则", claims_text)
        self.assertIn("技术领域规则", claims_text)
        self.assertIn("背景技术规则", claims_text)
        self.assertIn("references/cases/claims-format-standard.md", claims_text)
        self.assertNotIn("说明书摘要规则", claims_text)
        self.assertNotIn("具体实施方式规则", claims_text)

        self.assertIn("说明书摘要规则", full_text)
        self.assertIn("发明内容规则", full_text)
        self.assertIn("具体实施方式规则", full_text)

        self.assertIn("权要返修", revision_text)
        self.assertIn("全文返修", revision_text)
        self.assertIn("留痕返修", revision_text)
        # 署名不再默认 Juventude，改为每次由用户提供
        self.assertIn("作者名不设默认值", revision_text)
        self.assertNotIn("默认作者名为 `Juventude`", revision_text)

        self.assertIn("sectPr", docx_text)
        self.assertIn("header*.xml", docx_text)
        self.assertIn("headerReference", docx_text)
        self.assertIn("不得替代 `docx` skill 的执行层", docx_text)

        self.assertIn("附图自动生成规则", figures_text)
        self.assertIn("图 1", figures_text)
        self.assertIn("Visio", figures_text)
    def test_audit_reinforcements_are_preserved(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        rules_text = (SKILL_DIR / "rules.md").read_text(encoding="utf-8")
        global_text = (SKILL_DIR / "references" / "rules" / "global.md").read_text(encoding="utf-8")
        claims_text = (SKILL_DIR / "references" / "rules" / "claims.md").read_text(encoding="utf-8")
        full_text = (SKILL_DIR / "references" / "rules" / "full-draft.md").read_text(encoding="utf-8")
        revision_text = (SKILL_DIR / "references" / "rules" / "revision.md").read_text(encoding="utf-8")
        docx_text = (SKILL_DIR / "references" / "rules" / "docx-template.md").read_text(encoding="utf-8")
        figures_text = (SKILL_DIR / "references" / "rules" / "figures.md").read_text(encoding="utf-8")

        self.assertIn("单一出处原则", rules_text)
        self.assertIn("唯一权威出处是 `references/rules/global.md` G2", rules_text)
        self.assertIn("所有阶段默认强制读取", global_text)
        self.assertNotIn("当前唯一强制规则文件", global_text)
        self.assertEqual(global_text.count("## 0. 规则分级总表"), 1)

        self.assertIn("输入对象 → 关键处理机制 → 中间结果 → 参与最终结果", global_text)
        self.assertIn("训练阶段的历史输入", global_text)
        self.assertIn("使用阶段的当前输入", global_text)
        self.assertIn("交底书 DOCX 转出的带批注 Markdown 统一命名", global_text)

        self.assertIn("背景技术默认写成 2–3 个自然段", claims_text)
        self.assertIn("第一段", claims_text)
        self.assertIn("第二段", claims_text)
        self.assertIn("第三段", claims_text)
        self.assertIn("两个相对独立且均有技术贡献", claims_text)
        self.assertIn("功能模块式系统/装置权", claims_text)
        self.assertIn("不得展开技术效果、实施步骤或背景缺陷", claims_text)

        self.assertIn("只控制最终可见文本", skill_text)
        self.assertIn("不删除、不重建非当前阶段分节锚点", docx_text)
        self.assertIn("非当前阶段模板槽位", docx_text)
        self.assertIn("后续阶段内容不得因权要一稿被全局清空", docx_text)
        self.assertIn("权要一稿/二稿/三稿最终 DOCX 可见页眉仅为", docx_text)
        self.assertIn("仍保留包内全部 `word/header*.xml`", docx_text)
        self.assertIn("案例性术语清理", skill_text)
        self.assertIn("后续阶段槽位保留不删", skill_text)

        self.assertIn("全文终稿**目标**约 1.5–2 万字", full_text)
        self.assertIn("render_patent_figure.py", figures_text)
        self.assertIn("insert_figures_docx.py", figures_text)
        self.assertIn("子流程图默认不画", figures_text)

        self.assertIn("同类问题", revision_text)
        self.assertIn("原稿、返修稿和 `comments.xml`", revision_text)
        self.assertIn("对应阶段规则文件", skill_text)

    def test_second_reaudit_alignment_fixes_are_preserved(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        global_text = (SKILL_DIR / "references" / "rules" / "global.md").read_text(encoding="utf-8")
        claims_text = (SKILL_DIR / "references" / "rules" / "claims.md").read_text(encoding="utf-8")
        full_text = (SKILL_DIR / "references" / "rules" / "full-draft.md").read_text(encoding="utf-8")
        revision_text = (SKILL_DIR / "references" / "rules" / "revision.md").read_text(encoding="utf-8")
        docx_text = (SKILL_DIR / "references" / "rules" / "docx-template.md").read_text(encoding="utf-8")
        figures_text = (SKILL_DIR / "references" / "rules" / "figures.md").read_text(encoding="utf-8")

        self.assertNotIn("由用户自行编译后粘贴到 DOCX", global_text)
        self.assertIn("一律写成 **LaTeX 源码**", global_text)
        self.assertIn("LaTeX 源码", skill_text)
        self.assertIn("权利要求中不写公式", global_text)
        self.assertIn("解释每个参数", global_text)

        self.assertIn("构造保证", figures_text)
        self.assertIn("替换语义", figures_text)
        self.assertIn("重新生成", figures_text)
        self.assertIn("用户自备", figures_text)

        for row in [
            "| 全局 1级规则 | 所有块、所有阶段 |",
            "| 全局 2级规则 | 所有块、所有阶段 |",
            "| 局部 1级规则 | 某一块 |",
            "| 局部 2级规则 | 某一块 |",
        ]:
            self.assertEqual(global_text.count(row), 1, row)

        self.assertNotIn("方法 ≤7 条", claims_text)
        # 默认配比示例守卫: 必须与"总数=10"算术闭合。旧文案"方法权通常写到权要 7 左右"
        # + "权9 系统、权10 存储介质"只有 9 条(漏权8), 已改为方法权到权 8。
        self.assertIn("方法权通常写到权要 8", claims_text)
        self.assertNotIn("方法权通常写到权要 7", claims_text)
        self.assertIn("可扩展方法权数量或压缩系统", claims_text)
        self.assertIn("系统/装置独权默认采用计算机设备式写法", claims_text)
        self.assertIn("计算机可读存储介质权", claims_text)
        self.assertNotIn("系统/装置独权及其从权的结构", skill_text)

        self.assertIn("构建模型", global_text)
        self.assertIn("生成策略", global_text)
        self.assertIn("补偿信号", global_text)
        self.assertIn("进行评估", global_text)
        self.assertIn("输入参数 → 控制参数 → 初步控制信号", global_text)

        self.assertIn("交付类型：留痕稿/干净稿", skill_text)
        self.assertIn("默认不清除批注", revision_text)
        self.assertIn("最新已审权要 DOCX", skill_text)
        self.assertIn("候选不唯一", skill_text)
        self.assertIn("document-skills:docx", skill_text)
        self.assertIn("docx` skill", docx_text)

        self.assertIn("未来槽位锚点仍可定位", docx_text)
        self.assertIn("保留 XML 锚点", docx_text)
        self.assertIn("不得删除整节", docx_text)
        self.assertIn("不得全局清空 body", docx_text)

        self.assertIn("无方法独立权要", full_text)
        self.assertIn("主要保护主题", full_text)
        self.assertIn("附图说明一致", full_text)
        self.assertIn("句末标点", figures_text)
        self.assertIn("分号分句", figures_text)

    def test_dedup_and_segmented_write_rules_are_preserved(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        rules_text = (SKILL_DIR / "rules.md").read_text(encoding="utf-8")
        global_text = (SKILL_DIR / "references" / "rules" / "global.md").read_text(encoding="utf-8")
        full_text = (SKILL_DIR / "references" / "rules" / "full-draft.md").read_text(encoding="utf-8")
        docx_text = (SKILL_DIR / "references" / "rules" / "docx-template.md").read_text(encoding="utf-8")
        figures_text = (SKILL_DIR / "references" / "rules" / "figures.md").read_text(encoding="utf-8")

        # 单一出处原则与指针化
        self.assertIn("单一出处原则", rules_text)
        self.assertIn("仅适用于**权要一稿**和**全文一稿**", global_text)

        # 全文稿分块撰写法：Markdown 层 + DOCX 注入层（脚本化注入）
        self.assertIn("全文稿分块撰写法", full_text)
        self.assertIn("不得一次性生成全文长文", global_text)
        self.assertIn("inject_fulltext_docx.py", full_text)
        self.assertIn("一次性注入", full_text)
        self.assertIn("要求用户手动复制粘贴进 Word", full_text)
        self.assertIn("全文稿分块撰写法", skill_text)

        # G8-0 分节对照表与 G8-0b 标题格式
        self.assertIn("分节↔页眉↔正文内容对照表", docx_text)
        self.assertIn("G8-0b", docx_text)
        self.assertIn("顶格", docx_text)
        self.assertIn("<w:b/>", docx_text)
        self.assertIn("五个章节标题", docx_text)

        # 摘要附图唯一出处对齐 full-draft L5
        self.assertIn("full-draft.md` L5", figures_text)
        self.assertNotIn("摘要附图默认 = 图 1", figures_text)

        # 公式 LaTeX 源码交付
        self.assertIn("不带 `$`/`$$` 定界符", global_text)

        # 案件 docs/ 目录与 skill 自身 docs/ 消歧
        self.assertIn("docs/` 子目录", global_text)
        self.assertIn("技能文档目录", global_text)


    def test_extract_structure_and_cross_block_pass_on_conforming_draft(self):
        import json

        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(GOOD_FULL_DRAFT_MD, encoding="utf-8")

            result = run_script("extract_structure.py", "--md", md_path, "--stage", "full-draft")
            data = json.loads(result.stdout)
            self.assertTrue(data["extraction_ok"])
            claim1 = next(i for i in data["claims"]["items"] if i["num"] == 1)
            self.assertEqual(claim1["step_count"], 3)
            self.assertEqual(data["main_steps"]["ids"], ["S11", "S12", "S13"])
            self.assertEqual(data["figures"]["nums"], [1, 2])
            claim3 = next(i for i in data["claims"]["items"] if i["num"] == 3)
            self.assertEqual(claim3["range_refs"], [[1, 2]])

            result = run_script("check_cross_block.py", "--md", md_path, "--stage", "full-draft")
            data = json.loads(result.stdout)
            self.assertEqual(data["violation_count"], 0)

    def test_check_cross_block_catches_step_mismatch_and_bad_dependency(self):
        import json

        bad_md = GOOD_FULL_DRAFT_MD.replace(
            "在步骤S13中，根据中间结果生成控制信号。\n", ""
        ).replace("根据权利要求1所述", "根据权利要求9所述")
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(bad_md, encoding="utf-8")

            result = run_script_allow_fail(
                "check_cross_block.py", "--md", md_path, "--stage", "full-draft"
            )
            self.assertNotEqual(result.returncode, 0)
            data = json.loads(result.stdout)
            checks = {v["check"] for v in data["violations"]}
            self.assertIn("X1", checks)  # 权 1 分句 3 vs 主步骤 2
            self.assertIn("X3", checks)  # 依附了不存在的权要 9

    def test_check_hard_rules_judgment_sentence_pattern(self):
        import json

        def run_case(impl_body):
            md = "## 具体实施方式\n\n" + impl_body + "\n"
            with tempfile.TemporaryDirectory() as tmp:
                md_path = Path(tmp) / "全文稿.md"
                md_path.write_text(md, encoding="utf-8")
                result = run_script_allow_fail(
                    "check_hard_rules.py", "--md", md_path, "--stage", "full-draft", "--json"
                )
                data = json.loads(result.stdout)
                return [v for v in data["violations"]
                        if "判断" in v["message"] or "落单" in v["message"] or "半支" in v["message"]]

        # 合法: 规范句式二 / 单分支合并 / 名词性描述 / 白话"若是" 均不报 (规则 23 零误报边界)
        self.assertEqual(run_case(
            "判断所述残差是否超过预设阈值，若是，则触发告警；若否，则继续采集。\n\n"
            "若所述温度越限，则执行降载操作。\n\n"
            "所述阈值作为判断是否进入剧烈反应阶段的依据，用于判定是否出现卡顿。\n\n"
            "此时若是首次采集，需要初始化缓存。"
        ), [])
        # 违规: 判断步骤缺规定分支 / 句式一半支 / 若否落单
        self.assertEqual(len(run_case("在步骤S12中，判断所述数据是否有效，若有效则输出。")), 1)
        self.assertEqual(len(run_case("当判定满足触发条件时，则执行降载操作。")), 1)
        self.assertEqual(len(run_case("计算残差，若否，则继续采集。")), 1)

    def test_check_hard_rules_substep_generic_and_echo(self):
        import json

        def run_case(md, keys):
            with tempfile.TemporaryDirectory() as tmp:
                md_path = Path(tmp) / "全文稿.md"
                md_path.write_text(md, encoding="utf-8")
                result = run_script_allow_fail(
                    "check_hard_rules.py", "--md", md_path, "--stage", "full-draft", "--json"
                )
                data = json.loads(result.stdout)
                return [v for v in data["violations"] if any(k in v["message"] for k in keys)]

        # 规则 24: 子步骤编号残留 S111 抓到; 主步骤 S11 不报
        self.assertEqual(len(run_case(
            "## 具体实施方式\n\n在步骤S11中，处理数据，包括：步骤S111，读取数据。\n",
            ["子步骤编号"])), 1)
        self.assertEqual(run_case(
            "## 具体实施方式\n\n在步骤S11中，处理数据，包括：读取数据；解析数据。\n",
            ["子步骤编号"]), [])
        # 规则 25: 光杆泛词抓到; 带具体宾语不报
        self.assertEqual(len(run_case(
            "## 发明内容\n\n本方案提高效率，提升了系统稳定性。\n", ["光杆泛词"])), 2)
        self.assertEqual(run_case(
            "## 发明内容\n\n本方案提高了热源定位的准确性与响应速度。\n", ["光杆泛词"]), [])
        # 规则 26: 缺收口 / 双收口 / 呼应断裂 抓到; 逐字呼应放行
        self.assertEqual(len(run_case(
            "## 背景技术\n\n现有方法不好用。\n\n## 发明内容\n\n本发明以解决现有问题。\n",
            ["标准收口"])), 1)
        self.assertEqual(len(run_case(
            "## 背景技术\n\n方法甲，导致误报率高的问题。方法乙，导致漏报频发的问题。\n",
            ["必须唯一"])), 1)
        self.assertEqual(len(run_case(
            "## 背景技术\n\n方法甲不适配，导致误报率高的问题。\n\n"
            "## 发明内容\n\n本发明提供一种方法，以解决漏报频发的技术问题。\n",
            ["逐字包含"])), 1)
        self.assertEqual(run_case(
            "## 背景技术\n\n方法甲不适配，导致误报率高的问题。\n\n"
            "## 发明内容\n\n本发明提供一种方法，以解决现有技术无法自适应过滤、导致误报率高的问题。\n",
            ["标准收口", "必须唯一", "逐字包含", "以解决"]), [])

    def test_check_hard_rules_suspect_channel(self):
        import json

        md = (
            "## 具体实施方式\n\n"
            "在步骤S11中，采用数字调光方式调节PWM占空比。\n\n"
            "在步骤S12中，通过电流镜像电路进行模拟调流补偿，并将特征输入预设的状态识别模型，输出状态类别。\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(md, encoding="utf-8")
            result = run_script_allow_fail(
                "check_hard_rules.py", "--md", md_path, "--stage", "full-draft", "--json"
            )
            data = json.loads(result.stdout)
        # 规则 27 互斥体系共现 + 规则 28 模型缺维度 → suspect 线索, 不计 FAIL.
        # 断言按"本 fixture 触发的线索类型"判定, 不锁死总数: 新增探测器 (W08/W11/
        # W14/W16/W20/W26/W46) 各有独立 fixture, 本例只守历史两条线索仍在场.
        self.assertTrue(any("互斥技术体系" in s["message"] for s in all_suspects(data)))
        self.assertTrue(any("四维度" in s["message"] for s in all_suspects(data)))
        self.assertTrue(all("互斥" not in v["message"] and "四维度" not in v["message"]
                            for v in data["violations"]))
        # suspect 不进 exit code / 不计 FAIL
        self.assertNotIn("suspect", " ".join(v["message"] for v in data["violations"]))

    def test_check_cross_block_x7_dep_quote_verbatim(self):
        import json

        good_md = """## 权利要求书

1.一种测试控制方法，其特征在于，包括：获取输入数据；根据输入数据确定中间结果；根据中间结果生成控制信号。

2.根据权利要求1所述的测试控制方法，其特征在于，所述根据输入数据确定中间结果，包括：对输入数据滤波；确定中间结果。

## 技术领域

本发明涉及测试领域。

## 背景技术

现有技术存在问题。
"""

        def run_case(md_text):
            with tempfile.TemporaryDirectory() as tmp:
                md_path = Path(tmp) / "权要稿.md"
                md_path.write_text(md_text, encoding="utf-8")
                result = run_script_allow_fail(
                    "check_cross_block.py", "--md", md_path, "--stage", "claims-draft"
                )
                data = json.loads(result.stdout)
                return [v for v in data["violations"] if v["check"] == "X7"]

        # 规范引用 (仅去连接词/加"所述") 通过
        self.assertEqual(run_case(good_md), [])
        # 改动词 (根据→对): 引用句不再是权 1 原文连续子串, X7 FAIL
        bad_verb = good_md.replace(
            "所述根据输入数据确定中间结果，包括", "所述对输入数据确定中间结果，包括"
        )
        self.assertEqual(len(run_case(bad_verb)), 1)
        # 引用句内附加原句没有的限定, X7 FAIL
        bad_extra = good_md.replace(
            "所述根据输入数据确定中间结果，包括",
            "所述根据输入数据确定中间结果，其中所述中间结果为预设模型输出，包括",
        )
        self.assertEqual(len(run_case(bad_extra)), 1)

    def test_extract_structure_fails_hard_on_nonstandard_writing(self):
        import json

        bad_md = "## 权利要求书\n\n1.一种测试方法，包括：步骤甲；步骤乙。\n"
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "权要稿.md"
            md_path.write_text(bad_md, encoding="utf-8")

            result = run_script_allow_fail(
                "extract_structure.py", "--md", md_path, "--stage", "claims-draft"
            )
            self.assertEqual(result.returncode, 2)
            data = json.loads(result.stdout)
            self.assertFalse(data["extraction_ok"])
            self.assertTrue(any("其特征在于" in e for e in data["extraction_errors"]))

    def test_split_workflow_full_draft_uses_claims_md(self):
        import json

        # 真实工作流: 全文稿.md 不含权利要求书 (冻结在权要稿.md)
        idx = GOOD_FULL_DRAFT_MD.index("## 技术领域")
        claims_part = GOOD_FULL_DRAFT_MD[:idx]
        full_part = GOOD_FULL_DRAFT_MD[idx:]
        with tempfile.TemporaryDirectory() as tmp:
            claims_path = Path(tmp) / "权要稿.md"
            full_path = Path(tmp) / "全文稿.md"
            claims_path.write_text(claims_part, encoding="utf-8")
            full_path.write_text(full_part, encoding="utf-8")

            # 不传 --claims-md: 报结构错误并提示传入
            result = run_script_allow_fail(
                "check_cross_block.py", "--md", full_path, "--stage", "full-draft"
            )
            self.assertNotEqual(result.returncode, 0)
            data = json.loads(result.stdout)
            self.assertTrue(any("--claims-md" in e for e in data["extraction_errors"]))

            # 传 --claims-md: 全部通过
            result = run_script(
                "check_cross_block.py", "--md", full_path, "--stage", "full-draft",
                "--claims-md", claims_path,
            )
            data = json.loads(result.stdout)
            self.assertEqual(data["violation_count"], 0)
            self.assertEqual(data["structure"]["claims_source"], "claims_md")

            # 分离式全文稿跑第一类脚本: 不误报权要缺失/章节缺失
            result = run_script_allow_fail(
                "check_hard_rules.py", "--md", full_path, "--stage", "full-draft", "--json"
            )
            hard = json.loads(result.stdout)
            msgs = " ".join(v["message"] for v in hard["violations"])
            self.assertNotIn("未找到权要 1", msgs)
            self.assertNotIn("缺少章节: 权利要求书", msgs)

    def test_check_cross_block_flags_merged_step_lead(self):
        import json

        merged_md = GOOD_FULL_DRAFT_MD.replace(
            "在步骤S12中，根据输入数据确定中间结果。\n\n在步骤S13中，根据中间结果生成控制信号。\n",
            "在步骤S12至步骤S13中，根据输入数据确定中间结果，再根据中间结果生成控制信号。\n",
        )
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(merged_md, encoding="utf-8")

            result = run_script_allow_fail(
                "check_cross_block.py", "--md", md_path, "--stage", "full-draft"
            )
            self.assertNotEqual(result.returncode, 0)
            data = json.loads(result.stdout)
            checks = {v["check"] for v in data["violations"]}
            self.assertIn("X4", checks)   # 合并展开被点名
            self.assertNotIn("X1", checks)  # 覆盖数 3 == 分句数 3, 不误报数量

    def test_skill_wires_cross_block_check_into_gates(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")

        self.assertIn("scripts/check_cross_block.py", skill_text)
        self.assertIn("structure_check_result", skill_text)
        # 多路审查: SKILL.md 编排四个 auditor 契约
        for name in ("claims-auditor", "content-auditor", "impl-auditor", "global-auditor"):
            self.assertIn(name, skill_text)
        # 各契约消费 structure_check_result 且引用脚本
        for fname in ("claims-auditor.md", "content-auditor.md", "impl-auditor.md", "global-auditor.md"):
            text = (SKILL_DIR / "agents" / fname).read_text(encoding="utf-8")
            self.assertIn("structure_check_result", text)
            self.assertIn("check_hard_rules", text)
        # 专审契约有必审语义项清单
        for fname in ("claims-auditor.md", "content-auditor.md", "impl-auditor.md"):
            text = (SKILL_DIR / "agents" / fname).read_text(encoding="utf-8")
            self.assertIn("必审语义项", text)

    def test_signature_not_defaulted_to_juventude(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        revision_text = (SKILL_DIR / "references" / "rules" / "revision.md").read_text(encoding="utf-8")

        # 署名不设默认值，每次由用户提供
        self.assertIn("不设默认值", skill_text)
        self.assertIn("<用户指定署名>", skill_text)
        # 正文流程不再把 Juventude 作为默认署名硬写
        self.assertNotIn("默认 `Juventude`", skill_text)
        self.assertNotIn("署 Juventude", skill_text)
        self.assertNotIn("默认作者名为 `Juventude`", revision_text)

    def test_innovation_points_driven_by_disclosure_annotations(self):
        global_text = (SKILL_DIR / "references" / "rules" / "global.md").read_text(encoding="utf-8")
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")

        # 创新点以批注圈定为准，不自行判断
        self.assertIn("创新点以交底书批注为准", global_text)
        self.assertIn("不自行另判", skill_text)
        # disclosure-analyst 已废弃：事实提纲由主 agent 深读时亲自产出，批注驱动条款落在 SKILL.md 主流程
        self.assertFalse((SKILL_DIR / "agents" / "disclosure-analyst.md").exists())
        self.assertIn("批注圈定的创新点", skill_text)
        self.assertIn("优先审查要求（批注）", skill_text)

    def test_highlight_is_not_a_workflow_rule(self):
        global_text = (SKILL_DIR / "references" / "rules" / "global.md").read_text(encoding="utf-8")

        self.assertNotIn("高亮", global_text)
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        self.assertNotIn("高亮", skill_text)

    def test_delivery_filename_uses_writer_directory(self):
        global_text = (SKILL_DIR / "references" / "rules" / "global.md").read_text(encoding="utf-8")

        self.assertIn("撰写者目录白名单：`夏晓贝/`（经验档案 xxb）、`王培元/`（经验档案 wpy）", global_text)
        # 现行 G1-1: 对外交付用全称命名, 禁止简写与状态后缀
        self.assertIn("`案件号-稿次-作者-发明题目全称.docx`", global_text)
        self.assertIn("不得使用 `<撰写者>-<稿次>.docx` 简写", global_text)
        self.assertIn("不得加 `-留痕` 等状态后缀", global_text)
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        self.assertIn("案件号-权要1稿-作者-发明题目全称.docx", skill_text)

    def test_a_paradigm_anchor_matches_extractor(self):
        """L8-0 展开段入口句锚点与 extract_structure.py 的主步骤正则互相匹配(防规则句面与脚本漂移)。"""
        full_draft_text = (SKILL_DIR / "references" / "rules" / "full-draft.md").read_text(encoding="utf-8")
        self.assertIn("在步骤Sxx中，〔复述权1第xx分句原文〕，包括：", full_draft_text)
        extractor_src = (SKILL_DIR / "scripts" / "extract_structure.py").read_text(encoding="utf-8")
        self.assertIn("在步骤S", extractor_src)

    def test_slimming_regression_guards(self):
        """三项瘦身 + 增量复核的回归保护：防止改动回退成整篇传规则/全量重审。"""
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        scoring_text = (SKILL_DIR / "references" / "rules" / "scoring.md").read_text(encoding="utf-8")

        # 1) scoring 摘录已实体化为静态文件并按路径传入，不整篇传卡
        self.assertIn("不整篇传入任何 auditor", scoring_text)
        for route in ("claims", "content", "impl", "global"):
            self.assertTrue(
                (SKILL_DIR / "references" / "rules" / f"scoring-{route}.md").exists(),
                f"scoring-{route}.md 静态摘录缺失",
            )
        self.assertIn("scoring_excerpt_path", skill_text)
        self.assertNotIn("scoring_rules_content", skill_text)
        self.assertNotIn("scoring_excerpt_content", skill_text)

        # 2) 规则一律传路径：四路契约用 *_path 字段且白名单放行契约点名路径
        for name in ("claims-auditor", "content-auditor", "impl-auditor", "global-auditor"):
            contract = (SKILL_DIR / "agents" / f"{name}.md").read_text(encoding="utf-8")
            self.assertIn("scoring_excerpt_path", contract, name)
            self.assertNotIn("_content` |", contract, f"{name} 仍有内容内联字段")
            # 3) 增量复核：四路契约均支持 reaudit_context
            self.assertIn("reaudit_context", contract, name)
        self.assertIn("增量复核模式", skill_text)
        self.assertNotIn("恒重调", skill_text)

        # 4) facts.md 由主 agent 深读时亲自产出
        self.assertIn("docs/facts.md", skill_text)

        # 5) 闸门命令带 --md；权要基准回写条款存在
        self.assertNotIn("check_hard_rules.py <", skill_text)
        self.assertIn("回写更新 `docs/权要稿.md`", skill_text)
        self.assertIn("权要联动回写", skill_text)


    def test_check_env_runs_stdlib_only_and_reports_deps(self):
        import json

        # 自检器必须只用标准库 (否则"检查依赖的工具自己缺依赖"会死锁):
        # 用空 site (-S) + 清空 PYTHONPATH 剥离第三方包路径来近似"纯净解释器".
        import os

        env = dict(os.environ)
        env.pop("PYTHONPATH", None)
        proc = subprocess.run(
            [sys.executable, "-S", str(SCRIPTS_DIR / "check_env.py"), "--json"],
            text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            env=env, check=False,
        )
        # 脚本本身不能因为 import 第三方失败而崩溃: returncode 只应是
        # 0 (必需项齐) 或 缺失必需项数; -S 剥离 site 会让 pip 类必需项不可见.
        # 不锁具体数值 (必需项清单会演进, 如新增 Pillow), 断言 rc == 报告的缺失数.
        data = json.loads(proc.stdout)  # 崩溃则这里抛 JSONDecodeError / stdout 为空
        missing_required = sum(
            1 for c in data["checks"] if c["required"] and c["ok"] is False
        )
        self.assertEqual(
            proc.returncode, missing_required,
            msg=f"check_env 异常崩溃或退出码失真: {proc.stderr[-300:]}",
        )
        names = {c["name"] for c in data["checks"]}
        self.assertIn("python-docx", names)
        self.assertIn("pandoc", names)
        # 分类正确: python-docx 属可自动装的 pip 类, pandoc 属可选
        docx_check = next(c for c in data["checks"] if c["name"] == "python-docx")
        self.assertEqual(docx_check["category"], "pip")
        self.assertTrue(docx_check["auto_installable"])
        pandoc_check = next(c for c in data["checks"] if c["name"] == "pandoc")
        self.assertFalse(pandoc_check["required"])  # pandoc 为可选

    def test_skill_wires_env_check_into_onboarding(self):
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        self.assertIn("接入前环境自检", skill_text)
        self.assertIn("scripts/check_env.py", skill_text)
        self.assertTrue((SKILL_DIR / "docs" / "install.md").exists())

    def test_check_env_includes_docx_plugin_check(self):
        import json

        # 单宿主 (claude): docx 插件检测项恒在, python-docx 同为必需 pip 项
        r = run_script_allow_fail("check_env.py", "--json")
        d = json.loads(r.stdout)
        names = {c["name"] for c in d["checks"]}
        self.assertTrue(any("document-skills:docx" in n for n in names))
        self.assertTrue(any(c["name"] == "python-docx" for c in d["checks"]))

    def test_single_host_no_codex_residue(self):
        # 2026-07-30 起 codex 分支删除, 仅维护 claude 单宿主; 跨宿主适配层不得回流
        self.assertFalse((SKILL_DIR / "docs" / "porting.md").exists())
        self.assertFalse((SKILL_DIR / "agents" / "openai.yaml").exists())
        skill_text = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        self.assertNotIn("Codex", skill_text)
        for fname in ("global-auditor.md", "claims-auditor.md", "content-auditor.md", "impl-auditor.md"):
            auditor_text = (SKILL_DIR / "agents" / fname).read_text(encoding="utf-8")
            self.assertNotIn("跨宿主", auditor_text)
            self.assertNotIn("Codex", auditor_text)


class FigureScriptsSmokeTest(unittest.TestCase):
    """render_patent_figure.py + insert_figures_docx.py 冒烟：权要稿 → PNG → 注入 XML."""

    CLAIMS_MD = (
        "# 测试权要稿\n\n## 权利要求书\n\n"
        "1.一种测试数据处理方法，其特征在于，包括：\n\n"
        "采集测试输入数据，得到输入序列；\n\n"
        "对所述输入序列进行特征提取，得到特征向量；\n\n"
        "根据所述特征向量生成测试处理指令。\n\n"
        "2.根据权利要求1所述的测试数据处理方法，其特征在于，包括：省略。\n"
    )

    MINIMAL_SECT = (
        "    <w:p>\n      <w:pPr>\n        <w:sectPr>\n"
        '          <w:pgSz w:w="11906" w:h="16838"/>\n'
        "        </w:sectPr>\n      </w:pPr>\n    </w:p>\n"
    )

    MINIMAL_DOC = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">\n'
        "  <w:body>\n"
        + MINIMAL_SECT * 4
        + "    <w:sectPr>\n"
        '      <w:pgSz w:w="11906" w:h="16838"/>\n'
        "    </w:sectPr>\n"
        "  </w:body>\n"
        "</w:document>\n"
    )

    MINIMAL_RELS = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" '
        'Target="styles.xml"/></Relationships>\n'
    )

    MINIMAL_CT = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="xml" ContentType="application/xml"/></Types>\n'
    )

    def test_render_then_insert_roundtrip(self):
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            claims = tmp / "权要稿.md"
            claims.write_text(self.CLAIMS_MD, encoding="utf-8")
            png = tmp / "figure-1.png"
            result = run_script(
                "render_patent_figure.py", "--claims-md", claims, "--output", png
            )
            self.assertTrue(png.exists())
            self.assertEqual(png.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")
            self.assertIn("S11", result.stdout)

            unpacked = tmp / "unpacked"
            (unpacked / "word" / "_rels").mkdir(parents=True)
            (unpacked / "word" / "document.xml").write_text(
                self.MINIMAL_DOC, encoding="utf-8"
            )
            (unpacked / "word" / "_rels" / "document.xml.rels").write_text(
                self.MINIMAL_RELS, encoding="utf-8"
            )
            (unpacked / "[Content_Types].xml").write_text(
                self.MINIMAL_CT, encoding="utf-8"
            )
            run_script("insert_figures_docx.py", unpacked, "--png", png)

            doc = (unpacked / "word" / "document.xml").read_text(encoding="utf-8")
            self.assertEqual(doc.count("<w:drawing>"), 2)  # 分节2 + 分节5
            self.assertIn(">图1</w:t>", doc)  # 分节5 图题
            rels = (unpacked / "word" / "_rels" / "document.xml.rels").read_text(
                encoding="utf-8"
            )
            self.assertIn("media/image1.png", rels)
            self.assertTrue((unpacked / "word" / "media" / "image1.png").exists())
            ct = (unpacked / "[Content_Types].xml").read_text(encoding="utf-8")
            self.assertIn('Extension="png"', ct)

    def test_render_rejects_single_clause_claim(self):
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            claims = tmp / "权要稿.md"
            claims.write_text(
                "## 权利要求书\n\n1.一种测试系统，其特征在于，包括：处理器。\n",
                encoding="utf-8",
            )
            result = run_script_allow_fail(
                "render_patent_figure.py",
                "--claims-md", claims, "--output", tmp / "figure-1.png",
            )
            self.assertNotEqual(result.returncode, 0)


class InjectFulltextSmokeTest(unittest.TestCase):
    """inject_fulltext_docx.py + verify_docx_injection.py 冒烟：全文稿 md → 模板 DOCX 注入。

    用真实模板（assets/docx/专利撰写模板.docx）unpack 出 5 sectPr 骨架 + 一份含块公式
    与行内 $...$ 公式的全文稿.md fixture，真跑注入与校验脚本，断言段落数、oMath 数、
    无 $ 残留、sectPr 仍为 5。@skipUnless pandoc。
    """

    FULLTEXT_MD_FIXTURE = (
        "## 说明书摘要\n\n"
        "本发明涉及测试领域，公开了一种测试方法，获取输入数据；根据输入数据确定中间结果；"
        "根据中间结果生成控制信号。本发明实现了测试效果。\n\n"
        "## 摘要附图\n\n图1\n\n"
        "## 发明内容\n\n"
        "本发明提供一种测试方法，包括：\n"
        "获取输入数据；\n"
        "根据所述输入数据确定所述中间结果；\n"
        "根据所述中间结果生成所述控制信号。\n\n"
        "## 附图说明\n\n图1为测试方法流程示意图。\n\n"
        "## 具体实施方式\n\n"
        "在步骤S11中，获取输入数据，包括：读取传感器数据；对传感器数据滤波。\n\n"
        "其中，核函数的表达式为：\n\n"
        r"K(x, y) = \exp(-\gamma \|x - y\|^{2})" "\n\n"
        "式中，$K(x, y)$ 为核函数输出，$\\gamma$ 为核参数，$\\|x - y\\|$ 为欧氏距离，"
        "本实施例中取 $\\gamma=0.5$。\n\n"
        "综上所述，本发明公开了一种测试方法。本发明实现了测试效果。\n\n"
        "本发明第二实施例提供了一种测试系统，包括存储器、处理器及存储在存储器上并可在"
        "处理器上运行的计算机程序，所述处理器执行所述计算机程序时实现上述测试方法的步骤。\n\n"
        "需要说明的是，本发明实施例提供的一种测试系统用于执行上述实施例的一种测试方法的"
        "所有流程步骤，两者的工作原理和有益效果一一对应，因而不再赘述。\n\n"
        "以上所述的具体实施例，对本发明的目的、技术方案和有益效果进行了进一步的详细说明，"
        "应当理解，以上所述仅为本发明的具体实施例而已，并不用于限定本发明的保护范围。\n"
    )

    def _docx_skill_dir(self):
        """探测 document-skills:docx skill 目录，避免硬编码 commit hash 路径。

        优先级：环境变量 DOCX_SKILL_DIR → ~/.claude/plugins/cache 下最新 document-skills
        commit → 兜底硬编码本机当前 hash。他机/CI 用 DOCX_SKILL_DIR 覆盖即可。
        """
        env = os.environ.get("DOCX_SKILL_DIR")
        if env and Path(env).is_dir():
            return Path(env)
        cache = Path.home() / ".claude" / "plugins" / "cache" / "anthropic-agent-skills"
        if cache.is_dir():
            # 按 mtime 取最新: commit hash 是十六进制, 反字典序 ≠ 最新
            cands = sorted(cache.glob("document-skills/*/skills/docx"),
                           key=lambda q: q.stat().st_mtime, reverse=True)
            if cands:
                return cands[0]
        fallback = (cache / "document-skills" / "690f15cac7f7" / "skills" / "docx")
        return fallback

    def _unpack_template(self, dest: Path):
        tpl = SKILL_DIR / "assets" / "docx" / "专利撰写模板.docx"
        subprocess.run(
            [sys.executable, str(self._docx_skill_dir() / "scripts" / "office" / "unpack.py"),
             str(tpl), str(dest)],
            check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        )

    def _pack(self, unpacked: Path, out: Path):
        subprocess.run(
            [sys.executable, str(self._docx_skill_dir() / "scripts" / "office" / "pack.py"),
             str(unpacked), str(out)],
            check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
        )

    @unittest.skipUnless(PANDOC_AVAILABLE, "pandoc 不可用, 跳过全文稿注入测试")
    def test_inject_fulltext_then_verify(self):
        with tempfile.TemporaryDirectory() as td:
            tmp = Path(td)
            md = tmp / "全文稿.md"
            md.write_text(self.FULLTEXT_MD_FIXTURE, encoding="utf-8")
            unpacked = tmp / "unpack"
            self._unpack_template(unpacked)

            run_script("inject_fulltext_docx.py", unpacked, "--md", md)

            doc = (unpacked / "word" / "document.xml").read_text(encoding="utf-8")
            # 骨架未损
            self.assertEqual(len(re.findall(r"<w:sectPr[ >]", doc)), 5)
            # 块公式编译（1 条 LaTeX 独立成段）
            self.assertEqual(len(re.findall(r"<m:oMathPara\b", doc)), 1)
            # 行内 $...$ 实例数对账（$ 数必为偶数）
            om_inline = len(re.findall(r"<m:oMath\b", doc)) - len(re.findall(r"<m:oMathPara\b", doc))
            md_text = md.read_text(encoding="utf-8")
            self.assertEqual(om_inline, md_text.count("$") // 2)
            # 无 $ 残留、无裸 LaTeX
            all_t = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", doc))
            self.assertNotIn("$", all_t)
            self.assertNotIn(r"\exp", all_t)
            # 章节标题齐全（加粗）
            for title in ("发明内容", "附图说明", "具体实施方式"):
                # 在含该标题文本的段内查 <w:b/>
                self.assertTrue(re.search(
                    r"<w:p\b[^>]*>(?:(?!</w:p>).)*?<w:b/>(?:(?!</w:p>).)*?>"
                    + title + r"</w:t>", doc, re.S))
            # 占位套话段被替换为 md 真实套话
            self.assertNotIn("......", all_t)
            self.assertIn("综上所述", all_t)

            # pack 后跑 verify_docx_injection（图1 未注入时 drawing 项会 FAIL，单独验证公式/骨架项）
            out = tmp / "out.docx"
            self._pack(unpacked, out)
            result = run_script_allow_fail(
                "verify_docx_injection.py", out, "--md", md,
            )
            # 公式/骨架/标题/套话项应 PASS；图1 drawing 项预期 FAIL（本测试未注图）
            self.assertIn("[PASS] 块公式 oMathPara=1", result.stdout)
            self.assertIn("[PASS] 行内 oMath=", result.stdout)
            self.assertIn("[PASS] sectPr=5", result.stdout)
            self.assertIn("[PASS] $残留=0", result.stdout)
            self.assertIn("[FAIL] 图1 drawing=0", result.stdout)  # 未注图，预期失败


class RuleAnchorGuardTest(unittest.TestCase):
    """verify_rule_anchors.py 的四类防护 (审计实测的四个洞).

    全部在 /tmp 隔离副本上做破坏实验, 不触碰仓库文件。
    """

    def _fake_skill(self, tmp):
        """在 tmp 下搭一份 scripts/ + agents/ + references/rules/ 的副本."""
        root = Path(tmp)
        shutil.copytree(SCRIPTS_DIR, root / "scripts")
        shutil.copytree(SKILL_DIR / "agents", root / "agents")
        (root / "references").mkdir()
        shutil.copytree(SKILL_DIR / "references" / "rules", root / "references" / "rules")
        return root

    def _run_guard(self, root, *extra):
        return subprocess.run(
            [sys.executable, str(root / "scripts" / "verify_rule_anchors.py"), *extra],
            text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False,
        )

    def test_baseline_clean_state_passes(self):
        """未破坏时 12 条全 PASS、exit=0 (基线与实际一致)."""
        result = self._run_guard(SKILL_DIR)
        self.assertEqual(result.returncode, 0, f"基线与实际不一致:\n{result.stderr[-800:]}")
        self.assertIn("失败 0", result.stderr)

    def test_detects_interval_collapse(self):
        """洞1: 区间内出现同名结束锚点 → 提取塌缩, 必须 FAIL 而非"PASS 正文 1 行".

        sed 的 /start/,/end/ 首次命中即闭合。旧逻辑只判 `if not body`(空才 FAIL),
        L8 区间从 77 行塌到 1 行(L8-1/L8-2/L8-3 全丢)仍报 PASS —— 规则静默消失。
        """
        with tempfile.TemporaryDirectory() as tmp:
            root = self._fake_skill(tmp)
            fd = root / "references" / "rules" / "full-draft.md"
            lines = fd.read_text(encoding="utf-8").splitlines()
            for i, l in enumerate(lines):
                if l.startswith("## L8. "):
                    lines.insert(i + 3, "## 全文阶段自检重点")
                    break
            fd.write_text("\n".join(lines), encoding="utf-8")
            result = self._run_guard(root)
        self.assertNotEqual(result.returncode, 0, "区间塌缩未被检出")
        self.assertIn("区间塌缩", result.stderr)

    def test_double_quoted_anchor_is_parsed(self):
        """洞2: 契约用双引号书写 sed 时锚点不得静默隐身 (仍应解析到 12 条)."""
        with tempfile.TemporaryDirectory() as tmp:
            root = self._fake_skill(tmp)
            ia = root / "agents" / "impl-auditor.md"
            s = ia.read_text(encoding="utf-8")
            old = "sed -n '/^## L8\\. /,/^## 全文阶段自检重点/p'"
            self.assertIn(old, s, "fixture 前提失效: 未找到单引号形态锚点")
            ia.write_text(s.replace(old, 'sed -n "/^## L8\\. /,/^## 全文阶段自检重点/p"', 1),
                          encoding="utf-8")
            result = self._run_guard(root)
        self.assertIn("共 12 条", result.stderr, "双引号锚点被漏解析")
        self.assertEqual(result.returncode, 0, f"双引号形态误报:\n{result.stderr[-500:]}")

    def test_detects_deleted_anchor_via_count(self):
        """洞3: 误删一条锚点 → 条数校验必须亮红 (旧逻辑 12→11 仍报 0 失败)."""
        with tempfile.TemporaryDirectory() as tmp:
            root = self._fake_skill(tmp)
            ia = root / "agents" / "impl-auditor.md"
            s = ia.read_text(encoding="utf-8")
            old = "sed -n '/^### G6-1 /,/^### G6-2 /p'"
            self.assertIn(old, s, "fixture 前提失效")
            ia.write_text(s.replace(old, "(锚点已删)", 1), encoding="utf-8")
            result = self._run_guard(root)
        self.assertNotEqual(result.returncode, 0, "误删锚点未被检出")
        self.assertIn("期望 12 条", result.stderr)

    def test_update_baseline_refuses_on_collapse(self):
        """洞4: --update-baseline 在塌缩状态下必须拒绝, 不把错误固化为基线."""
        with tempfile.TemporaryDirectory() as tmp:
            root = self._fake_skill(tmp)
            fd = root / "references" / "rules" / "full-draft.md"
            lines = fd.read_text(encoding="utf-8").splitlines()
            for i, l in enumerate(lines):
                if l.startswith("## L8. "):
                    lines.insert(i + 3, "## 全文阶段自检重点")
                    break
            fd.write_text("\n".join(lines), encoding="utf-8")
            before = (root / "scripts" / "verify_rule_anchors.py").read_text(encoding="utf-8")
            result = self._run_guard(root, "--update-baseline")
            after = (root / "scripts" / "verify_rule_anchors.py").read_text(encoding="utf-8")
        self.assertEqual(result.returncode, 1, "塌缩时 --update-baseline 未拒绝")
        self.assertIn("拒绝更新", result.stderr)
        self.assertEqual(before, after, "塌缩状态下基线被改写 → 错误已固化")


class CrossBlockFlatStripTest(unittest.TestCase):
    """emit 层剥离 `flat` 省 token, 且不得让 X7 静默失效."""

    def test_flat_stripped_from_emitted_json(self):
        """emit 的 JSON 不含 flat (下游零消费, 实测占比 22.4%)."""
        import json
        md = GOOD_FULL_DRAFT_MD
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "权要稿.md"
            md_path.write_text(md, encoding="utf-8")
            result = run_script_allow_fail(
                "check_cross_block.py", "--md", md_path, "--stage", "claims-draft",
            )
            self.assertNotIn('"flat"', result.stdout, "emit 的 JSON 仍含 flat, 省 token 失效")
            data = json.loads(result.stdout)  # 仍是合法 JSON
            self.assertIn("structure", data)

    def test_x7_reports_input_incomplete_when_flat_missing(self):
        """回读已剥离 flat 的 structure 时, X7 必须报输入不完整而非静默 PASS.

        X7 靠 flat 做逐字子串匹配。若 emit 剥离后用户再经 --structure 回读,
        flat 缺失会让真违规也报 PASS —— 与"合规"输出完全同形、无法分辨。
        故必须显式报错。
        """
        import json
        md = GOOD_FULL_DRAFT_MD
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "权要稿.md"
            md_path.write_text(md, encoding="utf-8")
            first = run_script_allow_fail(
                "check_cross_block.py", "--md", md_path, "--stage", "claims-draft",
            )
            structure = json.loads(first.stdout)["structure"]
            # 该 md 有从权 (权2 依附权1), 满足触发条件
            self.assertTrue(any(it.get("dependent") for it in structure["claims"]["items"]),
                            "fixture 需含至少一条从权才能验 X7")
            struct_path = Path(tmp) / "structure.json"
            struct_path.write_text(json.dumps(structure, ensure_ascii=False), encoding="utf-8")
            reread = run_script_allow_fail(
                "check_cross_block.py", "--structure", struct_path,
            )
            data = json.loads(reread.stdout)
        x7 = [v for v in data["violations"] if v.get("check") == "X7"]
        self.assertEqual(len(x7), 1, "缺 flat 时 X7 未报输入不完整 → 静默失效风险复发")
        self.assertIn("缺 `flat` 字段", x7[0]["message"])

    def test_extraction_failure_does_not_crash_strip(self):
        """抽取失败时 claims 为 None (非缺键), 剥离逻辑不得崩."""
        import json
        idx = GOOD_FULL_DRAFT_MD.index("## 技术领域")
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(GOOD_FULL_DRAFT_MD[idx:], encoding="utf-8")
            result = run_script_allow_fail(
                "check_cross_block.py", "--md", md_path, "--stage", "full-draft",
            )
        # 必须是合法 JSON (未崩), 且报出缺 --claims-md
        data = json.loads(result.stdout)
        self.assertTrue(any("--claims-md" in e for e in data["extraction_errors"]))


class SuspectManifestRoutingTest(unittest.TestCase):
    """suspect manifest 唯一路由校验 (设计稿 §2.4/§5.5).

    每个脚本产出的 suspect 必须: ①带稳定 suspect_id; ②有唯一合法 owner;
    ③恰好出现在 manifest 的一个 owner 分组下; ④不进 exit code.
    W46 是唯一按章节互斥路由者: L8→impl, 其他说明书块→global, 未知章节转
    violation 而非广播多路.
    """

    VALID_OWNERS = {"global", "impl", "content", "claims"}

    # 七个探测器全部触发的 fixture (取自 approval_items 真实批注 anchor 语料)
    TRIGGER_MD = (
        "## 权利要求书\n\n"
        "1.一种充电桩电源的故障诊断预警方法，其特征在于，包括：获取数据。\n\n"
        "## 发明内容\n\n"
        "本发明不判断设备类型。\n\n"
        "## 具体实施方式\n\n"
        "下面将结合本发明实施例中的附图进行描述。本实施方式中的充电桩电源包括交流输入侧的"
        "功率因数校正电路和连接于直流母线的直流变换电路，功率因数校正电路用于整流。\n\n"
        "在步骤S11中，获取数据，包括：采集电压。\n\n"
        "当原始电压数据达到预设切换电压，且原始电流数据由恒定状态转为下降状态时，"
        "确定充电桩电源处于恒流恒压切换阶段。\n\n"
        "允许偏差范围和稳定变化限值由同型号健康充电桩的历史切换数据预先标定。\n\n"
        "需要说明的是，健康充电桩是指经绝缘检测、输出电压精度检测均满足设备技术要求的充电桩。\n\n"
        "按照功率因数校正电路状态数据在前、直流变换电路状态数据在后的顺序沿特征维度拼接，"
        "得到运行状态数据集。\n\n"
        "波动超限值 = 波动指标 - 阈值\n\n"
        "当判定不存在故障显性风险时，不生成预警指令。\n"
    )

    def _run(self, md_text, *extra):
        import json
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(md_text, encoding="utf-8")
            result = run_script_allow_fail(
                "check_hard_rules.py", "--md", md_path,
                "--stage", "full-draft", "--json", *extra,
            )
            return json.loads(result.stdout), result.returncode

    def test_every_suspect_has_unique_owner(self):
        """每条 suspect 都有稳定 id 与唯一合法 owner, manifest 分组无遗漏无重复."""
        data, _ = self._run(self.TRIGGER_MD)
        new = [s for s in all_suspects(data) if s.get("suspect_id")]
        self.assertTrue(new, "七探测器 fixture 未产出任何带 suspect_id 的线索")
        for s in new:
            self.assertTrue(s["suspect_id"], "suspect 缺 suspect_id")
            self.assertIn(s["owner"], self.VALID_OWNERS, f"非法 owner: {s['owner']}")
            self.assertTrue(s["section"], f"{s['suspect_id']} 缺 section")
        manifest = data["suspect_manifest"]
        self.assertEqual(
            sum(len(v) for v in manifest.values()), data["suspect_count"],
            "manifest 分组总数与 suspect 总数不等 (存在遗漏或重复归属)",
        )

    def test_seven_detectors_all_fire(self):
        """七个新探测器在触发 fixture 上各出一条线索."""
        data, _ = self._run(self.TRIGGER_MD)
        ids = {s["suspect_id"] for s in all_suspects(data) if s.get("suspect_id")}
        for expected in ("S-W08-impl-opening", "S-W11-state-bounds",
                         "S-W14-calibration", "S-W16-vague-def",
                         "S-W20-concat-schema", "S-W26-trivial-formula",
                         "S-W46-negative-only"):
            self.assertIn(expected, ids, f"探测器未触发: {expected}")

    def test_w46_section_routing_mutually_exclusive(self):
        """W46: L8 归 impl, 其他说明书块归 global, 同一线索不跨路重复."""
        data, _ = self._run(self.TRIGGER_MD)
        w46 = [s for s in all_suspects(data) if s.get("suspect_id") == "S-W46-negative-only"]
        routes = {s["section"]: s["owner"] for s in w46}
        self.assertEqual(routes.get("L8"), "impl", "L8 的 W46 线索应归 impl")
        self.assertEqual(routes.get("L6"), "global", "L6 的 W46 线索应归 global")
        manifest = data["suspect_manifest"]
        impl_w46 = [x for x in manifest["impl"] if x["suspect_id"] == "S-W46-negative-only"]
        global_w46 = [x for x in manifest["global"] if x["suspect_id"] == "S-W46-negative-only"]
        self.assertTrue(all(x["section"] == "L8" for x in impl_w46))
        self.assertTrue(all(x["section"] != "L8" for x in global_w46))

    def test_suspects_not_in_exit_code(self):
        """只含 suspect 触发、无 hard 违规时 exit code 为 0."""
        md = (
            "## 说明书摘要\n\n本发明涉及测试领域，公开了一种测试方法。\n\n"
            "## 摘要附图\n\n图1\n\n"
            "## 发明内容\n\n本发明提供一种测试方法。\n\n"
            "## 附图说明\n\n图1为测试方法流程示意图。\n\n"
            "## 具体实施方式\n\n"
            "如图1所示，本发明实施例提供的一种测试方法，包括步骤S11至步骤S11：\n\n"
            "在步骤S11中，获取数据。\n\n"
            "差值 = 甲量 - 乙量\n\n"
            "综上所述，本发明公开了一种测试方法。本发明通过测试，实现了测试效果。\n\n"
            "本发明第二实施例提供了一种测试系统，包括存储器、处理器及存储在存储器上并可在"
            "处理器上运行的计算机程序，所述处理器执行所述程序时实现如上述所述的一种测试方法。\n\n"
            "需要说明的是，本发明实施例提供的一种测试系统用于执行上述实施例的一种测试方法的"
            "所有流程步骤，两者的工作原理和有益效果一一对应，因而不再赘述。\n\n"
            "以上所述的具体实施例，并不用于限定本发明的保护范围。\n"
        )
        data, code = self._run(md, "--invention-name", "一种测试方法")
        self.assertTrue(any(s.get("suspect_id") == "S-W26-trivial-formula"
                            for s in all_suspects(data)), "应触发初等算术线索")
        self.assertEqual(data["violation_count"], 0,
                         f"该 fixture 不应有 hard 违规: {[v['message'][:40] for v in data['violations']]}")
        self.assertEqual(code, 0, "suspect 不得计入 exit code")

    def test_manifest_owners_declared_in_contracts(self):
        """每个 suspect_id 在其 owner 的 auditor 契约中有强制消费条款 (设计稿 §5.5)."""
        data, _ = self._run(self.TRIGGER_MD)
        contracts = {
            owner: (SKILL_DIR / "agents" / f"{owner}-auditor.md").read_text(encoding="utf-8")
            for owner in self.VALID_OWNERS
        }
        for s in all_suspects(data):
            sid = s.get("suspect_id")
            if not sid:
                continue
            self.assertIn(
                sid, contracts[s["owner"]],
                f"{sid} 未在 agents/{s['owner']}-auditor.md 中声明强制消费",
            )


class InventionNameSlotTest(unittest.TestCase):
    """规则 31 正式题名槽位 (W35): 漏后缀/重复拼接/输入不足三态."""

    # A 类槽位 (须写正式题名全称): 摘要句 + 技术领域句 + 发明内容目的句.
    # B 类槽位 (收尾"综上所述"段) 固定写方法名, 不参与本规则判定.
    BASE = (
        "## 说明书摘要\n\n"
        "本发明涉及测试领域，公开了一种{name}，获取数据。\n\n"
        "## 技术领域\n\n"
        "本发明涉及测试领域，具体涉及一种{name}。\n\n"
        "## 权利要求书\n\n"
        "1.一种充电桩电源的故障诊断预警方法，其特征在于，包括：获取数据。\n\n"
        "9.一种充电桩电源的故障诊断预警系统，包括存储器、处理器。\n\n"
        "## 发明内容\n\n"
        "本发明的目的在于提供一种{name}，旨在解决现有问题。\n\n"
        "## 具体实施方式\n\n"
        "综上所述，本发明公开了一种充电桩电源的故障诊断预警方法。本发明实现了预警效果。\n"
    )
    FULL_NAME = "一种充电桩电源的故障诊断预警方法及系统"

    def _run(self, name_in_md, *extra):
        import json
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(self.BASE.format(name=name_in_md), encoding="utf-8")
            result = run_script_allow_fail(
                "check_hard_rules.py", "--md", md_path,
                "--stage", "full-draft", "--json", *extra,
            )
            return json.loads(result.stdout)

    def test_missing_suffix_detected(self):
        """A 类槽位漏"及系统"三字 → 三处各报一次 (摘要/技术领域/目的句)."""
        data = self._run("充电桩电源的故障诊断预警方法", "--invention-name", self.FULL_NAME)
        hits = [v for v in data["violations"] if v["rule_id"] == "G8-0b"]
        self.assertEqual(len(hits), 3, f"三个 A 类槽位应各报一次漏后缀, 实得 {len(hits)}")
        self.assertTrue(all("漏后缀" in v["message"] for v in hits))

    def test_verbatim_match_passes(self):
        """A 类逐字一致 → 不报."""
        data = self._run("充电桩电源的故障诊断预警方法及系统", "--invention-name", self.FULL_NAME)
        self.assertEqual([v for v in data["violations"] if v["rule_id"] == "G8-0b"], [])

    def test_closing_summary_is_b_class_not_flagged(self):
        """回归: 收尾"综上所述，本发明公开了一种〔权1保护主题〕"属 B 类, 写方法名不得报违规.

        真实稿 X2607084 曾被误报: 该段句式的唯一出处是 docx-template.md G8-2 与
        full-draft.md L8-1 收尾条 (骨架为"〔发明名称〕方法"分体式), 在此写入含
        "及系统"的正式全称反而产出"……方法及系统方法"式重复拼接.
        """
        # A 类三处全部写全称, 收尾段写方法名 (BASE 固定如此) → 应 0 违规
        data = self._run("充电桩电源的故障诊断预警方法及系统", "--invention-name", self.FULL_NAME)
        hits = [v for v in data["violations"] if v["rule_id"] == "G8-0b"]
        self.assertEqual(hits, [], f"收尾段写方法名被误判为漏后缀: {[v['message'][:60] for v in hits]}")
        # 反向确认: 收尾段确实在稿中且只写方法名
        md = self.BASE.format(name=self.FULL_NAME)
        self.assertIn("综上所述，本发明公开了一种充电桩电源的故障诊断预警方法。", md)

    def test_missing_input_is_suspect_not_violation(self):
        """闸门死锁回归: 缺 --invention-name 走 suspect 通道, 不计 violation/exit code.

        该状态**无法通过改稿消除**(触发条件是参数没传, 不是稿子写错)。若计为
        violation, 按 SKILL.md"任一 FAIL 直接回修, 不进入阶段 B"的纪律, 主 agent
        会陷入死锁: 要么反复空转去改本来正确的题名句, 要么判定脚本误报而跳过整个
        闸门、丢失机械前置保护。故"配置缺失"与"稿件违规"必须分通道。
        """
        import json
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(self.BASE.format(name=self.FULL_NAME), encoding="utf-8")
            result = run_script_allow_fail(
                "check_hard_rules.py", "--md", md_path, "--stage", "full-draft", "--json",
            )
            data = json.loads(result.stdout)
        # 不得计入 violation
        self.assertEqual([v for v in data["violations"] if v["rule_id"] == "G8-0b"], [],
                         "缺参数被计为 violation → 闸门死锁复发")
        # 必须回一条配置缺失 suspect, 且明确区分于稿件违规
        hits = [s for s in all_suspects(data) if s.get("suspect_id") == "S-W35-name-input-missing"]
        self.assertEqual(len(hits), 1, "应恰好回一条配置缺失 suspect")
        self.assertIn("配置缺失", hits[0]["message"])
        self.assertIn("--invention-name", hits[0]["message"])
        self.assertIn("不得从权 1", hits[0]["message"])
        self.assertEqual(hits[0]["owner"], "global")

    def test_scan_excludes_non_delivery_sections(self):
        """越界回归: 只扫交付正文, 审查报告节内复述题名不得误判为槽位违规.

        审查报告在引述问题时常复述题名(如"摘要写'公开了一种xx方法'漏了'及系统'"),
        规则 31 若用 enumerate(lines) 扫全文会把这类引述当成正文槽位报违规。
        必须走 _get_scan_ranges(排除评分报告/审查报告/元数据/TODO)。
        """
        import json
        md = (
            "## 说明书摘要\n\n"
            "本发明涉及测试领域，公开了一种智能控制方法及系统，获取数据。\n\n"
            "## 发明内容\n\n"
            "本发明的目的在于提供一种智能控制方法及系统，旨在解决问题。\n\n"
            "## 审查报告\n\n"
            "本轮审查发现：说明书摘要写\"公开了一种智能控制方法\"，漏了\"及系统\"三字，需回修。\n\n"
            "## 具体实施方式\n\n"
            "在步骤S11中，获取数据。\n"
        )
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(md, encoding="utf-8")
            result = run_script_allow_fail(
                "check_hard_rules.py", "--md", md_path, "--stage", "full-draft",
                "--invention-name", "一种智能控制方法及系统", "--json",
            )
            data = json.loads(result.stdout)
        hits = [v for v in data["violations"] if v["rule_id"] == "G8-0b"]
        self.assertEqual(hits, [],
                         f"审查报告节内的题名引述被误判: {[v['location'] for v in hits]}")

    def test_skill_md_commands_pass_invention_name(self):
        """SKILL.md 的 full-draft 命令行必须带 --invention-name (调用方与脚本必需参数一致).

        脚本新增必需参数时若忘改调用方, 规则 31 会对所有案件静默失效
        (永远走"缺参数"分支, A 类槽位从此不校验)。本测试锁住两者同步。
        """
        skill = (SKILL_DIR / "SKILL.md").read_text(encoding="utf-8")
        # 取所有含 check_hard_rules.py 且 stage=full-draft 的命令行片段
        cmds = [seg for seg in re.findall(r"`[^`]*check_hard_rules\.py[^`]*`", skill)
                if "full-draft" in seg]
        self.assertTrue(cmds, "SKILL.md 未找到 full-draft 的 check_hard_rules 命令行")
        missing = [c for c in cmds if "--invention-name" not in c]
        self.assertEqual(
            missing, [],
            f"以下 full-draft 命令行缺 --invention-name (会导致 G8-0b 静默失效): {missing}",
        )


class SpecNoClaimsWordingTest(unittest.TestCase):
    """规则 29 说明书禁权要体例 (W18/W29/W30): 违规/合规/权要书内合法例外."""

    def _run(self, md_text):
        import json
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "全文稿.md"
            md_path.write_text(md_text, encoding="utf-8")
            result = run_script_allow_fail(
                "check_hard_rules.py", "--md", md_path, "--stage", "full-draft", "--json",
            )
            return json.loads(result.stdout)

    def test_claims_wording_in_spec_fails(self):
        """说明书出现权要体例措辞 → 逐处报 (取 W18/W29/W30 真实 anchor)."""
        md = (
            "## 权利要求书\n\n1.一种测试方法，其特征在于，包括：获取数据。\n\n"
            "## 具体实施方式\n\n"
            "多维空间欧氏距离属于公知距离度量，不在权利要求中增加其常规计算展开。\n\n"
            "所述计算机程序被处理器执行时，实现权利要求1至8任一项所述的方法的步骤。\n\n"
            "实际部署形式不影响处理器执行权利要求1至8任一项所述方法步骤。\n"
        )
        data = self._run(md)
        hits = [v for v in data["violations"] if "权要体例" in v["message"]]
        self.assertEqual(len(hits), 3, "三处权要体例措辞应各报一次")

    def test_claims_section_itself_exempt(self):
        """权利要求书章节内的"其特征在于"是法定体例 → 不报 (最易误报的合法例外)."""
        md = (
            "## 权利要求书\n\n"
            "1.一种测试方法，其特征在于，包括：获取数据。\n\n"
            "3.一种测试系统，包括存储器、处理器，其特征在于，"
            "所述处理器执行程序时实现所述的一种测试方法。\n\n"
            "## 具体实施方式\n\n"
            "本发明第二实施例提供了一种测试系统，包括：存储器、处理器及存储在存储器上并可在"
            "处理器上运行的计算机程序，所述处理器执行所述程序时实现如上述所述的一种测试方法。\n"
        )
        data = self._run(md)
        self.assertEqual([v for v in data["violations"] if "权要体例" in v["message"]], [])


if __name__ == "__main__":
    unittest.main()
